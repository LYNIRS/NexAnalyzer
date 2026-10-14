"""Raman-axis calibration against the laser line (Local 5.7.2)."""

import io
from io import BytesIO

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from core.io.report_settings import load_correct_raman_axis, save_correct_raman_axis
from modules.spectra.io.results_csv import export_master_csv
from modules.spectra.io.results_excel import TechniqueResults, build_sample_results_xlsx
from modules.spectra.models.peak import FitResult, FittedPeak
from modules.spectra.models.spectrum import (
    LaserCalibration,
    ProcessingSettings,
    SpectrumData,
    SpectrumFile,
)
from modules.spectra.processing import axis_calibration
from modules.spectra.processing.axis_calibration import (
    DEADBAND,
    axis_note,
    calibrate_raman_session,
    calibrate_raman_spectra,
    find_laser_line,
)

# The IMEC-Nov maps' axis: -680 to 1069 cm⁻¹ in 0.87 cm⁻¹ steps.
X = np.arange(-680.6, 1069.4, 0.8717)


def _gauss(x, center, fwhm, height):
    return height * np.exp(-0.5 * ((x - center) / (fwhm / 2.3548)) ** 2)


def _map(laser=1.7, n=4, laser_fwhm=2.4, laser_height=200.0, x=X, seed=0,
         raman_peak=251.5):
    """A map of `n` spectra: a laser line at `laser`, a Raman peak sitting at
    `raman_peak` on the *true* axis (so at raman_peak + laser as read)."""
    rng = np.random.default_rng(seed)
    spectra = []
    for _ in range(n):
        y = 470.0 + rng.normal(0, 2.0, x.size)
        if laser is not None:
            y += _gauss(x, laser, laser_fwhm, laser_height)
        y += _gauss(x, raman_peak + (laser or 0.0), 5.5, 60.0)
        spectra.append(SpectrumData(X=x, Y=y))
    return spectra


class TestFindLaserLine:
    def test_finds_an_offset_line(self):
        spectra = _map(laser=1.7)
        line, fwhm = find_laser_line(X, np.mean([s.Y for s in spectra], axis=0))
        assert line == pytest.approx(1.7, abs=0.05)
        assert fwhm == pytest.approx(2.4, abs=0.2)

    def test_line_between_filter_shoulders_is_still_found(self):
        # The IMEC-Nov RM_6 shape: notch-filter shoulders ~6 cm⁻¹ either side.
        y = 490.0 + _gauss(X, -1.9, 2.4, 160.0)
        y += _gauss(X, -8.0, 3.0, 30.0) + _gauss(X, 5.0, 4.0, 60.0)
        line, _ = find_laser_line(X, y)
        assert line == pytest.approx(-1.9, abs=0.1)

    def test_a_taller_shoulder_past_the_limit_does_not_hide_the_line(self):
        # The QC_KRX261001 RM_9 shape: a filter shoulder beyond MAX_OFFSET
        # outgrows the line. Taking only the tallest point found the shoulder,
        # failed the offset test and left the map uncorrected.
        y = 520.0 + _gauss(X, 4.7, 2.2, 400.0) + _gauss(X, 11.0, 3.0, 480.0)
        line, _ = find_laser_line(X, y)
        assert line == pytest.approx(4.7, abs=0.15)

    def test_axis_that_starts_above_zero_has_no_line(self):
        x = np.linspace(100.0, 1000.0, 1000)
        assert find_laser_line(x, _gauss(x, 520.0, 4.0, 100.0)) is None

    def test_pl_wavelength_axis_has_no_line(self):
        x = np.linspace(600.0, 900.0, 1000)
        assert find_laser_line(x, _gauss(x, 770.0, 35.0, 500.0)) is None

    def test_flat_window_has_no_line(self):
        rng = np.random.default_rng(1)
        assert find_laser_line(X, 470.0 + rng.normal(0, 2.0, X.size)) is None

    def test_too_wide_a_line_is_not_a_laser(self):
        assert find_laser_line(X, 470.0 + _gauss(X, 1.0, 9.0, 200.0)) is None

    def test_offset_beyond_max_is_rejected(self):
        y = 470.0 + _gauss(X, axis_calibration.MAX_OFFSET + 5.0, 2.4, 200.0)
        assert find_laser_line(X, y) is None


class TestCalibrateRamanSpectra:
    def test_shifts_the_axis_so_the_laser_sits_at_zero(self):
        shifted, calibration = calibrate_raman_spectra(_map(laser=1.7))

        assert calibration.applied
        assert calibration.shift == pytest.approx(-1.7, abs=0.05)
        assert shifted[0].X == pytest.approx(X + calibration.shift)
        # Y untouched; the Raman peak now reads at its true position.
        line, _ = find_laser_line(shifted[0].X, np.mean([s.Y for s in shifted], axis=0))
        assert line == pytest.approx(0.0, abs=0.05)
        peak = shifted[0].X[np.argmax(np.where(np.abs(shifted[0].X - 251.5) < 10,
                                               shifted[0].Y, -np.inf))]
        assert peak == pytest.approx(251.5, abs=0.9)

    def test_every_spectrum_of_the_map_gets_the_same_shift(self):
        shifted, _ = calibrate_raman_spectra(_map(laser=-3.9, n=6))
        assert all(np.array_equal(s.X, shifted[0].X) for s in shifted)

    def test_offset_inside_the_deadband_is_measured_but_not_applied(self):
        spectra = _map(laser=DEADBAND / 2)
        shifted, calibration = calibrate_raman_spectra(spectra)

        assert not calibration.applied
        assert calibration.laser_line == pytest.approx(DEADBAND / 2, abs=0.05)
        assert shifted[0].X is spectra[0].X

    def test_disabled_still_measures_the_line(self):
        spectra = _map(laser=1.7)
        shifted, calibration = calibrate_raman_spectra(spectra, enabled=False)

        assert calibration.shift == 0.0
        assert calibration.laser_line == pytest.approx(1.7, abs=0.05)
        assert shifted[0].X is spectra[0].X

    def test_no_line_leaves_the_axis_as_read(self):
        x = np.linspace(100.0, 1000.0, 1000)
        spectra = [SpectrumData(X=x, Y=470.0 + _gauss(x, 520.0, 4.0, 100.0))]
        shifted, calibration = calibrate_raman_spectra(spectra)

        assert calibration == LaserCalibration(None, None, 0.0)
        assert shifted[0].X is x

    def test_empty_input(self):
        shifted, calibration = calibrate_raman_spectra([])
        assert shifted == [] and calibration.laser_line is None


class TestSession:
    """A load's files share an offset; a file without its own line takes theirs."""

    def _no_line(self, seed=9):
        return _map(laser=None, n=3, seed=seed)

    def test_a_file_without_a_line_takes_the_sessions_median(self):
        files = [_map(laser=4.7, n=3, seed=s) for s in range(5)] + [self._no_line()]
        results = calibrate_raman_session(files)

        own = [cal for _, cal in results[:5]]
        assert all(not c.from_session and c.shift == pytest.approx(-4.7, abs=0.05) for c in own)
        shifted, borrowed = results[5]
        assert borrowed.from_session and borrowed.laser_line is None
        assert borrowed.shift == pytest.approx(-4.7, abs=0.05)
        assert shifted[0].X == pytest.approx(X + borrowed.shift)

    def test_a_file_whose_line_disagrees_takes_the_sessions_median(self):
        # A shoulder accepted as the line in one map of five.
        files = [_map(laser=4.7, n=3, seed=s) for s in range(4)] + [_map(laser=9.8, n=3, seed=7)]
        _, odd = calibrate_raman_session(files)[4]

        assert odd.from_session
        assert odd.laser_line == pytest.approx(9.8, abs=0.05)  # still recorded
        assert odd.shift == pytest.approx(-4.7, abs=0.05)

    def test_no_majority_lends_nothing(self):
        files = [_map(laser=4.7, n=3, seed=s) for s in range(2)] + [self._no_line(s) for s in range(3)]
        results = calibrate_raman_session(files)

        assert all(not cal.from_session for _, cal in results)
        assert [cal.applied for _, cal in results] == [True, True, False, False, False]

    def test_off_borrows_nothing_either(self):
        files = [_map(laser=4.7, n=3, seed=s) for s in range(3)] + [self._no_line()]
        assert all(cal.shift == 0.0 for _, cal in calibrate_raman_session(files, enabled=False))

    def test_a_single_file_is_its_own_session(self):
        _, cal = calibrate_raman_session([self._no_line()])[0]
        assert cal == LaserCalibration(None, None, 0.0)

    def test_the_note_says_which_files_borrowed(self):
        note = axis_note([LaserCalibration(4.68, 2.5, -4.68),
                          LaserCalibration(None, None, -4.68, from_session=True)])
        assert note == ("Raman axis shifted to laser line: laser at +4.68 cm⁻¹ moved to 0; "
                        "1 file(s) without a matching line of their own took the "
                        "session's line at +4.68 cm⁻¹")

    def test_round_trips_through_a_project(self):
        cal = LaserCalibration(9.8, 3.1, -4.7, from_session=True)
        assert LaserCalibration.from_dict(cal.to_dict()) == cal
        assert LaserCalibration.from_dict({"laser_line": 1.0, "fwhm": 2.0, "shift": -1.0}).from_session is False


class TestAxisNote:
    def test_all_files_corrected(self):
        note = axis_note([LaserCalibration(1.65, 2.4, -1.65), LaserCalibration(1.70, 2.4, -1.70)])
        assert note == "Raman axis shifted to laser line: laser at +1.65 to +1.70 cm⁻¹ moved to 0"

    def test_some_files_corrected(self):
        note = axis_note([LaserCalibration(-1.9, 2.4, 1.9), LaserCalibration(None, None, 0.0)])
        assert "on 1 of 2 files" in note and "rest uncorrected" in note

    def test_found_but_not_applied(self):
        assert axis_note([LaserCalibration(0.1, 2.4, 0.0)]) == (
            "Raman axis uncorrected; laser line at +0.10 cm⁻¹")

    def test_nothing_found(self):
        assert axis_note([LaserCalibration(None, None, 0.0)]) == (
            "Raman axis uncorrected; no laser line found")

    def test_never_checked_and_empty(self):
        assert axis_note([None]) is None
        assert axis_note([]) is None


def _fitted_raman(calibration):
    data = _map(n=1)[0]
    component = _gauss(data.X, 251.5, 5.5, 60.0)
    return SpectrumFile(
        filename="RM_1__1.txt", mode="Raman",
        original_data=data, raw_data=data, processed_data=data,
        processing_settings=ProcessingSettings(),
        laser_calibration=calibration,
        fit_result=FitResult(
            success=True,
            fitted_peaks=[FittedPeak(
                label="E2g+A1g", center=251.5, center_stderr=0.1,
                area=350.0, area_stderr=5.0, width_fwhm=5.5, width_stderr=0.2,
                shape=0.5, component_curve=component, color="#1f77b4",
            )],
            total_fit_curve=component, residuals=np.zeros_like(component),
            chi_squared=1.0, r_squared=0.95, convergence_time=0.1,
        ),
    )


class TestRecorded:
    def test_spectrum_file_round_trips_its_calibration(self):
        spectrum = _fitted_raman(LaserCalibration(1.68, 2.45, -1.68))
        restored = SpectrumFile.from_dict(spectrum.to_dict())
        assert restored.laser_calibration == LaserCalibration(1.68, 2.45, -1.68)

    def test_a_project_saved_before_the_calibration_loads_as_unchecked(self):
        saved = _fitted_raman(None).to_dict()
        del saved["laser_calibration"]
        assert SpectrumFile.from_dict(saved).laser_calibration is None

    def test_reset_to_raw_keeps_the_corrected_axis(self):
        spectrum = _fitted_raman(LaserCalibration(1.68, 2.45, -1.68))
        spectrum.reset_to_raw()
        assert spectrum.laser_calibration.shift == -1.68

    def test_master_csv_carries_laser_line_and_shift(self):
        files = {
            "RM_1__1.txt": _fitted_raman(LaserCalibration(1.68, 2.45, -1.68)),
            "RM_2__1.txt": _fitted_raman(None),
        }
        df = pd.read_csv(io.StringIO(export_master_csv(files)))

        assert list(df.columns[6:8]) == ["Laser_Line_cm1", "Axis_Shift_cm1"]
        assert df.loc[0, "Laser_Line_cm1"] == pytest.approx(1.68)
        assert df.loc[0, "Axis_Shift_cm1"] == pytest.approx(-1.68)
        assert pd.isna(df.loc[1, "Laser_Line_cm1"]) and pd.isna(df.loc[1, "Axis_Shift_cm1"])

    def test_workbook_carries_the_axis_on_raman_only(self):
        spectrum = _fitted_raman(LaserCalibration(1.68, 2.45, -1.68))
        note = axis_note([spectrum.laser_calibration])
        wb = load_workbook(BytesIO(build_sample_results_xlsx(
            sample_name="IMC260805", material_name="WSe2_truncate", report_date="2026-10-13",
            techniques=[TechniqueResults(
                label="Raman", point_spectra=[(1, spectrum)],
                axis_note=note, axis_columns=True,
            )],
        )))

        header, row = [list(r) for r in wb["Raman"].iter_rows(values_only=True)][:2]
        assert header[-2:] == ["Laser_Line_cm1", "Axis_Shift_cm1"]
        assert row[-2:] == [pytest.approx(1.68), pytest.approx(-1.68)]
        summary = [v for r in wb["Summary"].iter_rows(values_only=True) for v in r if v]
        assert "Raman axis" in summary and note in summary

    def test_workbook_without_axis_columns_is_unchanged(self):
        wb = load_workbook(BytesIO(build_sample_results_xlsx(
            sample_name="S", material_name="M", report_date="2026-10-13",
            techniques=[TechniqueResults(label="Raman", point_spectra=[(1, _fitted_raman(None))])],
        )))
        header = next(wb["Raman"].iter_rows(values_only=True))
        assert header[-1] == "Convergence_Time_s"


class TestSampleBatch:
    """The QC Report's load path: calibrates Raman per file, never PL."""

    @pytest.fixture
    def batch(self, monkeypatch):
        from core.report.models import SampleScan
        from modules.spectra.processing import sample_batch

        maps = {"RM_1.txt": _map(laser=1.7, n=3), "RM_2.txt": _map(laser=1.65, n=3, seed=2),
                "PL_1.txt": [SpectrumData(X=np.linspace(600, 900, 500),
                                          Y=_gauss(np.linspace(600, 900, 500), 770, 35, 500))]}
        monkeypatch.setattr(sample_batch, "parse_spectrum_multi", lambda path: maps[path])
        # execute_auto_workflow needs Streamlit; the axis is set before it runs.
        monkeypatch.setattr(sample_batch, "execute_auto_workflow", lambda s, p: {"success": True})
        scan = SampleScan(folder="/f", sample_name="S",
                          raman_files={1: "RM_1.txt", 2: "RM_2.txt"}, pl_files={1: "PL_1.txt"})

        def run(**kwargs):
            return sample_batch.run_sample_batch(scan, object(), object(), **kwargs)
        return run

    def test_raman_files_are_shifted_and_recorded(self, batch):
        result = batch()

        assert sorted(result.raman_calibrations) == [1, 2]
        assert result.raman_calibrations[1].shift == pytest.approx(-1.7, abs=0.05)
        point, spectrum = result.raman_spectra[0]
        assert spectrum.laser_calibration is result.raman_calibrations[point]
        assert spectrum.original_data.X == pytest.approx(X + spectrum.laser_calibration.shift)
        assert all(s.laser_calibration is None for _, s in result.pl_spectra)

    def test_off_measures_but_does_not_shift(self, batch):
        result = batch(correct_raman_axis=False)

        _, spectrum = result.raman_spectra[0]
        assert spectrum.laser_calibration.laser_line == pytest.approx(1.7, abs=0.05)
        assert spectrum.original_data.X == pytest.approx(X)


class TestSetting:
    def test_defaults_on(self, tmp_path):
        assert load_correct_raman_axis(tmp_path / "missing.json") is True

    def test_round_trip(self, tmp_path):
        path = tmp_path / "report_settings.json"
        save_correct_raman_axis(False, path)
        assert load_correct_raman_axis(path) is False
        save_correct_raman_axis(True, path)
        assert load_correct_raman_axis(path) is True
